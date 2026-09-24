import os
import uuid
import threading
from datetime import datetime
from openpyxl import Workbook, load_workbook
from app.config import settings

_LOCK = threading.Lock()

COLUMNS = [
    "ticket_id",
    "customer_name",
    "phone",
    "email",
    "category",          # Billing / Complaint / Support / Sales
    "inquiry_details",
    "status",             # open / in_progress / closed / manual_review_required
    "priority",           # low / medium / high
    "ai_confidence",      # 0-1, filled in after the AI attempts a reply
    "ai_reply",            # text of the reply the AI sent (if any)
    "escalation_reason",   # why it was escalated, if applicable
    "created_at",
    "updated_at",
]


def _ensure_workbook():
    path = settings.tickets_path
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if not os.path.exists(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Tickets"
        ws.append(COLUMNS)
        wb.save(path)


def _priority_for_category(category: str) -> str:
    return {
        "Complaint": "high",
        "Billing": "medium",
        "Sales": "medium",
        "Support": "low",
    }.get(category, "medium")


def create_ticket(customer_name, phone, email, category, inquiry_details) -> dict:
    _ensure_workbook()
    with _LOCK:
        wb = load_workbook(settings.tickets_path)
        ws = wb["Tickets"]
        ticket_id = f"AE-{datetime.utcnow().strftime('%Y%m%d')}-{str(uuid.uuid4())[:6].upper()}"
        now = datetime.utcnow().isoformat(timespec="seconds") + "Z"
        row = {
            "ticket_id": ticket_id,
            "customer_name": customer_name,
            "phone": phone,
            "email": email,
            "category": category,
            "inquiry_details": inquiry_details,
            "status": "open",
            "priority": _priority_for_category(category),
            "ai_confidence": "",
            "ai_reply": "",
            "escalation_reason": "",
            "created_at": now,
            "updated_at": now,
        }
        ws.append([row[c] for c in COLUMNS])
        wb.save(settings.tickets_path)
        return row


def list_tickets(status: str | None = None) -> list[dict]:
    _ensure_workbook()
    with _LOCK:
        wb = load_workbook(settings.tickets_path)
        ws = wb["Tickets"]
        rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            record = dict(zip(COLUMNS, r))
            if status is None or record["status"] == status:
                rows.append(record)
        return rows


def get_ticket(ticket_id: str) -> dict | None:
    for t in list_tickets():
        if t["ticket_id"] == ticket_id:
            return t
    return None


def update_ticket(ticket_id: str, **fields) -> dict | None:
    _ensure_workbook()
    with _LOCK:
        wb = load_workbook(settings.tickets_path)
        ws = wb["Tickets"]
        id_col = COLUMNS.index("ticket_id") + 1
        for row in ws.iter_rows(min_row=2):
            if row[id_col - 1].value == ticket_id:
                for key, value in fields.items():
                    if key in COLUMNS:
                        col_idx = COLUMNS.index(key)
                        row[col_idx].value = value
                updated_idx = COLUMNS.index("updated_at")
                row[updated_idx].value = datetime.utcnow().isoformat(timespec="seconds") + "Z"
                wb.save(settings.tickets_path)
                return dict(zip(COLUMNS, [c.value for c in row]))
        return None
